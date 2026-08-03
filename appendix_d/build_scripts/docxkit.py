"""Document construction kit for Applied Business Analytics appendices.

Implements the stylesheet and construction rules of
`Chapter_Style_Specification.md` v1.0 (sections 1-4) directly, so that the
built file carries the book's styles without depending on a binary template
surviving transfer.  Every value below is traceable to a row of the spec's
master table; see `verify.py` for the checks that hold it there.
"""

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Twips
import re

ACCENT = "57068C"
CALLOUT_FILL = "F5F2F8"
CALLOUT_HAIRLINE = "C9BFD6"
CODE_FILL = "F2F2F2"
CODE_BORDER = "D6D6D6"
SECONDARY = "404040"
LINK = "0000FF"


def _el(tag, **attrs):
    if ":" not in tag:
        tag = "w:" + tag
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e


def _sub(parent, tag, **attrs):
    e = _el(tag, **attrs)
    parent.append(e)
    return e


# --------------------------------------------------------------------------
# style definitions: (styleId, name, basedOn, ppr-dict, rpr-dict)
# spacing values are twips; 20 twips = 1 pt.  font sizes are half-points.
# --------------------------------------------------------------------------

STYLES = [
    # id, name, basedOn, before, after, line, ppr-extra, rpr
    ("Heading1", "Heading 1", "Normal", 360, 120, 276,
     dict(keepNext=1, keepLines=1, outlineLvl=0),
     dict(b=1, sz=28)),
    ("Heading2", "Heading 2", "Normal", 280, 80, 276,
     dict(keepNext=1, keepLines=1, outlineLvl=1),
     dict(b=1, i=1, sz=24)),
    ("Heading3", "Heading 3", "Normal", 240, 60, 276,
     dict(keepNext=1, keepLines=1, outlineLvl=2),
     dict(b=1, sz=23)),
    ("BodyText", "Body Text", "Normal", 0, 0, 276, {}, dict(sz=24)),
    ("BodyTextIndent", "Body Text Indent", "Normal", 0, 0, 276,
     dict(indFirstLine=360), dict(sz=24)),
    ("ChapterTitle", "Chapter Title", "Normal", 0, 120, 276,
     dict(keepNext=1, jc="center", outlineLvl=0), dict(b=1, sz=34)),
    ("ChapterSubtitle", "Chapter Subtitle", "Normal", 0, 300, 276,
     dict(keepNext=1, jc="center"), dict(i=1, sz=26)),
    ("Byline", "Byline", "Normal", 0, 120, None,
     dict(jc="center", contextualSpacing=1), dict(sz=22, color=SECONDARY)),
    ("MetadataLabel", "Metadata Label", "Normal", 240, 40, None,
     dict(keepNext=1), dict(b=1, caps=1, sz=20, color=ACCENT)),
    ("MetadataBody", "Metadata Body", "Normal", 0, 80, 276, {}, dict(sz=22)),
    ("ListNumber", "List Number", "Normal", 40, 80, 276,
     dict(indLeft=720, indHanging=360, contextualSpacing=1), dict(sz=24)),
    ("ListBullet", "List Bullet", "Normal", 40, 80, 276,
     dict(indLeft=720, indHanging=360, contextualSpacing=1), dict(sz=24)),
    ("ReferenceEntry", "Reference Entry", "Normal", 0, 120, 276,
     dict(indHangingLeft=720), dict(sz=24)),
    ("TableCaption", "Table Caption", "Normal", 300, 80, 276,
     dict(keepNext=1), dict(sz=22)),
    ("TableHeading", "Table Heading", "Normal", 40, 40, None, {},
     dict(b=1, sz=21)),
    ("TableText", "Table Text", "Normal", 40, 40, None, {}, dict(sz=21)),
    ("CodeCaption", "Code Caption", "Normal", 300, 80, 276,
     dict(keepNext=1), dict(b=1, sz=22)),
    ("CodeBlock", "Code Block", "Normal", 60, 240, None,
     dict(keepLines=1, indLeft=180, indRight=180, shade=CODE_FILL,
          codeborder=1), dict(rFonts="Courier New", sz=19)),
    ("CalloutLabel", "Callout Label", "Normal", 160, 60, 276,
     dict(keepNext=1, keepLines=1, indLeft=200, indRight=200,
          shade=CALLOUT_FILL, calloutborder=1),
     dict(b=1, caps=1, sz=19, color=ACCENT, spacing=30)),
    ("CalloutTitle", "Callout Title", "Normal", 0, 80, 276,
     dict(keepNext=1, keepLines=1, indLeft=200, indRight=200,
          shade=CALLOUT_FILL, calloutborder=1), dict(b=1, sz=23)),
    ("CalloutBody", "Callout Body", "Normal", 0, 80, 276,
     dict(keepNext=1, keepLines=1, indLeft=200, indRight=200,
          shade=CALLOUT_FILL, calloutborder=1), dict(sz=22)),
    ("CalloutSource", "Callout Source", "Normal", 0, 140, 276,
     dict(keepLines=1, indLeft=200, indRight=200, shade=CALLOUT_FILL,
          calloutborder=1), dict(sz=20, color=SECONDARY)),
    ("Spacer", "Spacer", "Normal", 0, 0, None, {}, dict(sz=14)),
    ("Footer", "Footer", "Normal", 0, 120, None,
     dict(jc="center"), dict(sz=20, color=SECONDARY)),
]


def _callout_borders(pbdr):
    _sub(pbdr, "left", val="single", sz=18, space=8, color=ACCENT)
    for side in ("top", "bottom", "right"):
        _sub(pbdr, side, val="single", sz=4, space=6, color=CALLOUT_HAIRLINE)


def _code_borders(pbdr):
    for side in ("top", "left", "bottom", "right"):
        _sub(pbdr, side, val="single", sz=4, space=4, color=CODE_BORDER)


def _build_style(styles_el, sid, name, based, before, after, line, ppr, rpr):
    s = _sub(styles_el, "style", type="paragraph", styleId=sid)
    _sub(s, "name", val=name)
    _sub(s, "basedOn", val=based)
    _sub(s, "qFormat")
    p = _sub(s, "pPr")

    if ppr.get("keepNext"):
        _sub(p, "keepNext")
    if ppr.get("keepLines"):
        _sub(p, "keepLines")
    if "shade" in ppr:
        _sub(p, "shd", val="clear", color="auto", fill=ppr["shade"])
    if ppr.get("calloutborder"):
        _callout_borders(_sub(p, "pBdr"))
    if ppr.get("codeborder"):
        _code_borders(_sub(p, "pBdr"))

    ind = {}
    if "indFirstLine" in ppr:
        ind["firstLine"] = ppr["indFirstLine"]
    if "indLeft" in ppr:
        ind["left"] = ppr["indLeft"]
    if "indRight" in ppr:
        ind["right"] = ppr["indRight"]
    if "indHanging" in ppr:
        ind["hanging"] = ppr["indHanging"]
    if "indHangingLeft" in ppr:
        ind["left"] = ppr["indHangingLeft"]
        ind["hanging"] = ppr["indHangingLeft"]
    if ind:
        _sub(p, "ind", **ind)

    sp = dict(before=before, after=after)
    if line:
        sp["line"] = line
        sp["lineRule"] = "auto"
    else:
        sp["line"] = 240
        sp["lineRule"] = "auto"
    _sub(p, "spacing", **sp)

    if ppr.get("contextualSpacing"):
        _sub(p, "contextualSpacing")
    if "jc" in ppr:
        _sub(p, "jc", val=ppr["jc"])
    if "outlineLvl" in ppr:
        _sub(p, "outlineLvl", val=ppr["outlineLvl"])

    r = _sub(s, "rPr")
    if "rFonts" in rpr:
        _sub(r, "rFonts", ascii=rpr["rFonts"], hAnsi=rpr["rFonts"],
             cs=rpr["rFonts"])
    if rpr.get("b"):
        _sub(r, "b")
    if rpr.get("i"):
        _sub(r, "i")
    if rpr.get("caps"):
        _sub(r, "caps")
    if "color" in rpr:
        _sub(r, "color", val=rpr["color"])
    if "spacing" in rpr:
        _sub(r, "spacing", val=rpr["spacing"])
    if "sz" in rpr:
        _sub(r, "sz", val=rpr["sz"])
        _sub(r, "szCs", val=rpr["sz"])


def new_document():
    doc = Document()
    styles_el = doc.styles.element

    # -- document defaults -------------------------------------------------
    for child in list(styles_el):
        if child.tag == qn("w:docDefaults"):
            styles_el.remove(child)
    dd = _el("docDefaults")
    rpd = _sub(dd, "rPrDefault")
    rpr = _sub(rpd, "rPr")
    _sub(rpr, "rFonts", ascii="Times New Roman", hAnsi="Times New Roman",
         eastAsia="Times New Roman", cs="Times New Roman")
    _sub(rpr, "sz", val=24)
    _sub(rpr, "szCs", val=24)
    _sub(rpr, "lang", val="en-US")
    ppd = _sub(dd, "pPrDefault")
    ppr = _sub(ppd, "pPr")
    _sub(ppr, "spacing", after=120, line=276, lineRule="auto")
    styles_el.insert(0, dd)

    # -- Normal defines nothing beyond the defaults ------------------------
    for child in list(styles_el):
        if child.tag == qn("w:style"):
            styles_el.remove(child)

    normal = _sub(styles_el, "style", type="paragraph", styleId="Normal",
                  default="1")
    _sub(normal, "name", val="Normal")
    _sub(normal, "qFormat")

    for spec in STYLES:
        _build_style(styles_el, *spec)

    # -- character styles --------------------------------------------------
    for sid, name, rprops in (
        ("Hyperlink", "Hyperlink", dict(u="single", color=LINK)),
        ("Emphasis", "Emphasis", dict(i=1)),
    ):
        s = _sub(styles_el, "style", type="character", styleId=sid)
        _sub(s, "name", val=name)
        _sub(s, "qFormat")
        r = _sub(s, "rPr")
        if rprops.get("i"):
            _sub(r, "i")
        if "u" in rprops:
            _sub(r, "u", val=rprops["u"])
        if "color" in rprops:
            _sub(r, "color", val=rprops["color"])

    # -- page setup --------------------------------------------------------
    sec = doc.sections[0]
    sec.page_width = Twips(12240)
    sec.page_height = Twips(15840)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Twips(1440))
    sec.header_distance = Twips(720)
    sec.footer_distance = Twips(720)

    # -- footer: centred page number --------------------------------------
    footer_p = sec.footer.paragraphs[0]
    footer_p.style = doc.styles["Footer"]
    r = footer_p.add_run()
    fld1 = _el("fldChar", fldCharType="begin")
    instr = _el("instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = _el("fldChar", fldCharType="end")
    r._r.append(fld1)
    r._r.append(instr)
    r._r.append(fld2)

    # remove the empty first paragraph python-docx leaves behind
    body = doc.element.body
    for p in body.findall(qn("w:p")):
        body.remove(p)

    return doc


# --------------------------------------------------------------------------
# inline markup: *italic*, **bold**, [text](url), `code`
# --------------------------------------------------------------------------

NO_SMART = {"Code Block", "CodeBlock", "Code Caption", "CodeCaption"}


def smarten(text):
    """Typographic quotes and apostrophes, for everything but code."""
    out, open_q = [], True
    for ch in text:
        if ch == '"':
            out.append("\u201c" if open_q else "\u201d")
            open_q = not open_q
        elif ch == "'":
            out.append("\u2019")
        else:
            out.append(ch)
    return "".join(out)


_TOKEN = re.compile(
    r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|(?<!\\)\*[^*]+(?<!\\)\*|`[^`]+`)")


def add_runs(doc, para, text, smart=True):
    if smart:
        text = smarten(text)
    for piece in _TOKEN.split(text):
        if not piece:
            continue
        if piece.startswith("[") and "](" in piece:
            label, url = piece[1:-1].split("](", 1)
            _add_hyperlink(doc, para, label, url)
        elif piece.startswith("**") and piece.endswith("**"):
            para.add_run(piece[2:-2]).bold = True
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            r = para.add_run(piece[1:-1])
            r.style = doc.styles["Emphasis"]
        elif piece.startswith("`") and piece.endswith("`"):
            r = para.add_run(piece[1:-1])
            r.font.name = "Courier New"
            r.font.size = Twips(190)  # 9.5 pt
            rf = r._r.get_or_add_rPr().get_or_add_rFonts()
            rf.set(qn("w:cs"), "Courier New")
        else:
            para.add_run(piece.replace("\\*", "*"))


def _add_hyperlink(doc, para, label, url):
    part = para.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = _el("hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = _sub(run, "rPr")
    _sub(rpr, "rStyle", val="Hyperlink")
    t = _sub(run, "t")
    t.set(qn("xml:space"), "preserve")
    t.text = label
    link.append(run)
    para._p.append(link)


def para(doc, style, text=""):
    p = doc.add_paragraph(style=doc.styles[style])
    if text:
        add_runs(doc, p, text, smart=style not in NO_SMART)
    return p


# --------------------------------------------------------------------------
# tables
# --------------------------------------------------------------------------

def _tbl_borders(tblPr):
    b = _sub(tblPr, "tblBorders")
    _sub(b, "top", val="single", sz=8, space=0, color="000000")
    _sub(b, "bottom", val="single", sz=8, space=0, color="000000")
    for side in ("left", "right", "insideH", "insideV"):
        _sub(b, side, val="nil")


def add_table(doc, number, title, headers, rows, widths=None):
    """Two-paragraph caption above, then an APA-style data table."""
    para(doc, "TableCaption", f"Table {number}")
    para(doc, "TableCaption", f"*{title}*")

    t = doc.add_table(rows=1, cols=len(headers))
    tbl = t._tbl
    tblPr = tbl.tblPr
    for child in list(tblPr):
        if child.tag in (qn("w:tblBorders"), qn("w:tblW"),
                         qn("w:tblLayout"), qn("w:tblStyle")):
            tblPr.remove(child)
    _sub(tblPr, "tblW", w=9360, type="dxa")
    _sub(tblPr, "tblLayout", type="fixed")
    _tbl_borders(tblPr)
    marg = _sub(tblPr, "tblCellMar")
    _sub(marg, "top", w=80, type="dxa")
    _sub(marg, "bottom", w=80, type="dxa")
    _sub(marg, "left", w=100, type="dxa")
    _sub(marg, "right", w=100, type="dxa")

    if widths is None:
        widths = [round(9360 / len(headers))] * len(headers)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = _el("tblGrid")
    for w in widths:
        _sub(grid, "gridCol", w=w)
    tblPr.addnext(grid)

    hdr = t.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    _sub(trPr, "tblHeader")
    _sub(trPr, "cantSplit")
    for i, (cell, text) in enumerate(zip(hdr.cells, headers)):
        _set_cell(doc, cell, text, "TableHeading", widths[i], header=True)

    for row in rows:
        tr = t.add_row()
        for i, (cell, text) in enumerate(zip(tr.cells, row)):
            _set_cell(doc, cell, text, "TableText", widths[i])
    return t


def _set_cell(doc, cell, text, style, width, header=False):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcW"):
            tcPr.remove(child)
    _sub(tcPr, "tcW", w=width, type="dxa")
    if header:
        b = _sub(tcPr, "tcBorders")
        _sub(b, "bottom", val="single", sz=6, space=0, color="000000")
    p = cell.paragraphs[0]
    p.style = doc.styles[style]
    for child in list(p._p):
        if child.tag == qn("w:r"):
            p._p.remove(child)
    if text:
        add_runs(doc, p, text)


def callout(doc, label, title, body, source=None):
    para(doc, "CalloutLabel", label)
    if title:
        para(doc, "CalloutTitle", title)
    for b in body:
        para(doc, "CalloutBody", b)
    if source:
        para(doc, "CalloutSource", source)
    para(doc, "Spacer")
